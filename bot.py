import os
import re
import json
import requests
import telebot

from openpyxl import Workbook
from openpyxl.styles import PatternFill

# ---------- ADMIN SOZLAMALARI ----------

# Bu yerga O'ZINGIZNING Telegram ID'ingizni yozasiz!
# O'zingizni @userinfobot ga /start deb yozib, ID ni olasiz.
ADMIN_ID = 357556285

USERS_FILE = "users.json"   # Foydalanuvchilar ro'yxatini shu faylga yozamiz

# TOKEN берём из переменной окружения (Render → Environment → TOKEN)
TOKEN = os.getenv("TOKEN")
if not TOKEN:
    raise RuntimeError("Переменная окружения TOKEN не задана!")

bot = telebot.TeleBot(TOKEN, parse_mode="HTML")

API_URL = "https://www.imei.kg/api/phys/imei/status?imei={imei}"


# ---------- FOYDALANUVCHINI SAQLASH / ADMINGA JB ----------

def save_user(message):
    """
    Botdan foydalangan userlarni users.json ga yozib boramiz.
    """
    user = message.from_user
    uid = str(user.id)
    username = user.username or ""
    first_name = user.first_name or ""
    last_name = user.last_name or ""

    try:
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
    except FileNotFoundError:
        data = {}

    # Mavjud bo'lsa yangilaymiz, bo'lmasa qo'shamiz
    data[uid] = {
        "username": username,
        "first_name": first_name,
        "last_name": last_name,
    }

    with open(USERS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def notify_admin_imeis(message, imeis):
    """
    IMEI ishlatilganida adminga kim ishlatgani haqida JB yuborish.
    """
    if not ADMIN_ID:
        return  # ADMIN_ID qo'yilmagan bo'lsa, hech narsa qilmaymiz

    user = message.from_user
    chat = message.chat

    username = f"@{user.username}" if user.username else "—"
    full_name = " ".join(
        part for part in [user.first_name, user.last_name] if part
    ) or "—"

    imei_text = "\n".join(f"• {i}" for i in imeis)

    text = f"""
🟢 IMEI botdan foydalanildi

👤 User ID: <code>{user.id}</code>
🔗 Username: {username}
👤 Ism: {full_name}
💬 Chat ID: <code>{chat.id}</code>

📱 IMEI(lar):
{imei_text}
"""

    try:
        bot.send_message(ADMIN_ID, text.strip())
    except Exception:
        # Adminga jo'natishda xato bo'lsa, bot ishini to'xtatmaymiz
        pass


# --------- KG статус (SIM / срок регистрации) ---------

def classify_kg_status(status_code: str, reg_info: dict) -> dict:
    """
    Қирғизистон SIM ва рўйхатдан ўтиш бўйича қисқа JB + Excel ранги.

    РАНГЛАР:
      RED     → SIM не работала
      PURPLE  → истёк 30-дневный период
      BLUE    → срок до даты / просто UNREGISTERED
      GREEN   → REGISTERED
      GRAY    → неизвестно
    """

    # Регистрация бўлимидаги барча текстларни йиғамиз
    txt_parts = []
    for v in (reg_info or {}).values():
        if isinstance(v, str):
            txt_parts.append(v)
    txt = " ".join(txt_parts)
    txt_low = txt.lower()

    # 0) REGISTERED → GREEN
    if status_code == "REGISTERED":
        return {
            "kg_color": "GREEN",
            "kg_chat": "🇰🇬 KG: зарегистрирован, можно пользоваться. 🟢",
            "kg_excel": "Зарегистрирован, можно пользоваться",
        }

    # 1) SIM хали умуман тушмаган → RED
    if "не использовалось ни в одной из сетей мобильных операторов связи кыргызской республики" in txt_low:
        return {
            "kg_color": "RED",
            "kg_chat": "🇰🇬 KG: SIM в сетях КР ещё не работала. 🔴",
            "kg_excel": "SIM в сетях КР ещё не работала",
        }

    # 2) истёк 30-дневный период → PURPLE
    if ("истёк 30-дневный" in txt_low) or ("истек 30-дневный" in txt_low):
        return {
            "kg_color": "PURPLE",
            "kg_chat": "🇰🇬 KG: истёк 30-дневный период, нужна регистрация. 🟣",
            "kg_excel": "Истёк 30-дневный период, нужна регистрация",
        }

    # 3) срок регистрации до ДАТЫ → BLUE
    if "срок регистрации" in txt_low:
        m = re.search(r"до\s+(\d{2}\.\d{2}\.\d{4})", txt)
        if m:
            date = m.group(1)
            return {
                "kg_color": "BLUE",
                "kg_chat": f"🇰🇬 KG: до {date}, потом нужна регистрация. 🔵",
                "kg_excel": f"До {date}, потом нужна регистрация",
            }

    # 4) просто UNREGISTERED → BLUE
    if status_code == "UNREGISTERED":
        return {
            "kg_color": "BLUE",
            "kg_chat": "🇰🇬 KG: не зарегистрирован, нужна регистрация. 🔵",
            "kg_excel": "Не зарегистрирован, нужна регистрация",
        }

    # 5) Номаълум → GRAY
    return {
        "kg_color": "GRAY",
        "kg_chat": "🇰🇬 KG: статус не удалось определить. ⚪️",
        "kg_excel": "Статус не удалось определить",
    }


# --------- API и парсинг ответа ---------

def fetch_imei_data(imei: str) -> dict:
    try:
        r = requests.get(API_URL.format(imei=imei), timeout=10)
        return r.json()
    except Exception:
        return {"status": "ERROR", "message": "Request failed"}


def parse_imei(data: dict) -> dict:
    if data.get("status") != "SUCCESS":
        kg = {
            "kg_color": "GRAY",
            "kg_chat": "🇰🇬 KG: ошибка при запросе. ⚪️",
            "kg_excel": "Ошибка при запросе",
        }
        return {
            "model": "Неизвестно",
            "fullname": "Неизвестно",
            "status_code": "UNKNOWN",
            "status_text": "Ошибка или неверный IMEI",
            **kg,
        }

    d = data.get("data", {}) or {}

    model = (
        d.get("gsmaMarketingName")
        or d.get("standardisedFullName")
        or "Неизвестно"
    )
    fullname = d.get("standardisedFullName") or model

    reg = d.get("registrationStatus", {}) or {}
    status_code = (reg.get("status") or "").upper()

    if status_code == "REGISTERED":
        status_text = "Зарегистрирован"
    elif status_code == "UNREGISTERED":
        status_text = "Не зарегистрирован"
    else:
        status_code = "UNKNOWN"
        status_text = "Неизвестно"

    kg = classify_kg_status(status_code, reg)

    return {
        "model": model,
        "fullname": fullname,
        "status_code": status_code,
        "status_text": status_text,
        "kg_color": kg["kg_color"],
        "kg_chat": kg["kg_chat"],
        "kg_excel": kg["kg_excel"],
    }


def format_text_answer(imei: str, info: dict) -> str:
    model = info["model"]
    fullname = info["fullname"]
    status_code = info["status_code"]
    status_text = info["status_text"]
    kg_chat = info["kg_chat"]

    if status_code == "REGISTERED":
        emoji = "✅"
    elif status_code == "UNREGISTERED":
        emoji = "❌"
    else:
        emoji = "ℹ️"

    text = f"""
📱 {imei}

<pre><code>Полное название: {fullname}
IMEI: {imei}
Статус: {emoji} {status_text}
{kg_chat}</code></pre>
"""
    return text.strip()


# --------- Excel ---------

def create_excel_xlsx(imei_list, info_map) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "IMEI Report"

    ws.append(["Полное название", "IMEI", "Статус регистрации", "KG статус"])

    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red   = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    blue  = PatternFill(start_color="C6D9F1", end_color="C6D9F1", fill_type="solid")
    purple = PatternFill(start_color="E6B8F7", end_color="E6B8F7", fill_type="solid")  # фиолет
    gray  = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

    for imei in imei_list:
        info = info_map[imei]

        fullname = info["fullname"]
        status_text = info["status_text"]
        kg_excel = info["kg_excel"]
        kg_color = info["kg_color"]
        status_code = info["status_code"]

        row = ws.max_row + 1
        ws.append([fullname, imei, status_text, kg_excel])

        # Колонка C — общий статус
        if status_code == "REGISTERED":
            ws[f"C{row}"].fill = green
        elif status_code == "UNREGISTERED":
            ws[f"C{row}"].fill = red
        else:
            ws[f"C{row}"].fill = blue

        # Колонка D — KG статус
        if kg_color == "RED":
            ws[f"D{row}"].fill = red
        elif kg_color == "BLUE":
            ws[f"D{row}"].fill = blue
        elif kg_color == "GREEN":
            ws[f"D{row}"].fill = green
        elif kg_color == "PURPLE":
            ws[f"D{row}"].fill = purple
        elif kg_color == "GRAY":
            ws[f"D{row}"].fill = gray

    file_path = "imei_result.xlsx"
    wb.save(file_path)
    return file_path


# --------- Handlers ---------

@bot.message_handler(commands=["start"])
def start_handler(message):
    # Userni saqlaymiz
    save_user(message)

    bot.reply_to(
        message,
        "Здравствуйте! 👋\n\n"
        "Отправьте IMEI (один или несколько).\n"
        "Бот покажет статус регистрации и KG-статус (SIM/срок регистрации).\n\n"
        "Excel версия → команда /excel"
    )


@bot.message_handler(commands=["excel", "exel"])
def excel_handler(message):
    # Userni saqlaymiz
    save_user(message)

    text = message.text
    numbers = re.findall(r"\d{10,20}", text)
    imeis = [n.strip() for n in numbers]

    if not imeis:
        bot.reply_to(message, "❗ После команды укажите IMEI списком.")
        return

    # Adminga kim, qaysi IMEI bilan /excel ishlatganini yuboramiz
    notify_admin_imeis(message, imeis)

    info_map = {}
    for imei in imeis:
        data = fetch_imei_data(imei)
        info_map[imei] = parse_imei(data)

    path = create_excel_xlsx(imeis, info_map)
    with open(path, "rb") as f:
        bot.send_document(message.chat.id, f, caption="📄 Excel готов!")


@bot.message_handler(content_types=["text"])
def text_handler(message):
    # Userni saqlaymiz
    save_user(message)

    text = message.text

    if text.startswith("/"):
        return

    imeis = re.findall(r"\d{10,20}", text)
    if not imeis:
        bot.reply_to(message, "❗ Отправьте IMEI (можно несколько).")
        return

    # Adminga kim va qaysi IMEIlar bilan foydalanganini yuboramiz
    notify_admin_imeis(message, imeis)

    for imei in imeis:
        data = fetch_imei_data(imei)
        info = parse_imei(data)
        bot.send_message(message.chat.id, format_text_answer(imei, info))


print("Бот запущен...")
bot.infinity_polling(skip_pending=True)
